"""Audit trail: search, hash-chain verification and export (AT-01..AT-05)."""

from __future__ import annotations

from datetime import date, timedelta

from fastapi import APIRouter, Depends, Query, Response
from sqlalchemy import func, select

from app.api.common import CurrentUser, DbSession
from app.core.deps import require
from app.core.rbac import Permission
from app.db.types import utcnow
from app.domain.dates import ist_day_start
from app.models.audit import AuditLog
from app.services import audit

router = APIRouter(prefix="/audit", tags=["Audit trail"])


@router.get("", summary="Search the audit trail")
def search(
    db: DbSession,
    user: User = Depends(require(Permission.VIEW_AUDIT)),
    entity_type: str | None = None,
    entity_id: str | None = None,
    action: str | None = None,
    user_id: str | None = None,
    period_from: date | None = None,
    period_to: date | None = None,
    search: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
) -> dict:
    stmt = select(AuditLog)
    if entity_type:
        stmt = stmt.where(AuditLog.entity_type == entity_type)
    if entity_id:
        stmt = stmt.where(AuditLog.entity_id == entity_id)
    if action:
        stmt = stmt.where(AuditLog.action == action)
    if user_id:
        stmt = stmt.where(AuditLog.user_id == user_id)
    if period_from:
        stmt = stmt.where(AuditLog.ts_utc >= ist_day_start(period_from))
    if period_to:
        stmt = stmt.where(
            AuditLog.ts_utc < ist_day_start(period_to + timedelta(days=1))
        )
    if search:
        term = f"%{search}%"
        stmt = stmt.where(
            AuditLog.reason.ilike(term)
            | AuditLog.user_email.ilike(term)
            | AuditLog.entity_id.ilike(term)
        )

    total = db.execute(select(func.count()).select_from(stmt.subquery())).scalar_one()
    rows = db.execute(
        stmt.order_by(AuditLog.ts_utc.desc(), AuditLog.id.desc())
        .limit(page_size).offset((page - 1) * page_size)
    ).scalars().all()

    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "rows": [_serialise(r) for r in rows],
    }


@router.get("/facets", summary="Distinct entity types and actions for the filter bar")
def facets(
    db: DbSession, user: User = Depends(require(Permission.VIEW_AUDIT))
) -> dict:
    entities = db.execute(
        select(AuditLog.entity_type, func.count(AuditLog.id)).group_by(AuditLog.entity_type)
    ).all()
    actions = db.execute(
        select(AuditLog.action, func.count(AuditLog.id)).group_by(AuditLog.action)
    ).all()
    return {
        "entity_types": [{"value": e, "count": c} for e, c in sorted(entities)],
        "actions": [{"value": a, "count": c} for a, c in sorted(actions)],
    }


@router.get("/verify", summary="Re-walk the hash chain and confirm nothing was altered (AT-03)")
def verify(
    db: DbSession, user: User = Depends(require(Permission.VIEW_AUDIT))
) -> dict:
    return audit.verify_chain(db)


@router.get("/export", summary="Export the audit trail for an auditor")
def export(
    db: DbSession,
    user: User = Depends(require(Permission.VIEW_AUDIT)),
    period_from: date | None = None,
    period_to: date | None = None,
    fmt: str = Query("xlsx", pattern="^(xlsx|csv)$"),
) -> Response:
    stmt = select(AuditLog)
    if period_from:
        stmt = stmt.where(AuditLog.ts_utc >= ist_day_start(period_from))
    if period_to:
        stmt = stmt.where(
            AuditLog.ts_utc < ist_day_start(period_to + timedelta(days=1))
        )
    rows = list(db.execute(stmt.order_by(AuditLog.ts_utc.asc(), AuditLog.id.asc())).scalars().all())

    from app.services import reports

    if fmt == "csv":
        report = {
            "code": "AUDIT", "name": "Audit Trail",
            "columns": [{"key": k, "label": lbl} for k, lbl in AUDIT_COLUMNS],
            "rows": [_serialise(r) for r in rows], "totals": {}, "summary": {},
        }
        return Response(
            content=reports.to_csv(report), media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": 'attachment; filename="Audit_Trail.csv"'},
        )

    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Audit Trail"
    headers = [lbl for _, lbl in AUDIT_COLUMNS]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="0F2A44")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        payload = _serialise(row)
        sheet.append([_text(payload.get(key)) for key, _ in AUDIT_COLUMNS])

    for index, (key, lbl) in enumerate(AUDIT_COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = (
            min(max(len(lbl) + 4, 14), 50)
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    meta = workbook.create_sheet("Metadata")
    meta.append(["Generated", utcnow().isoformat()])
    meta.append(["Rows", len(rows)])
    meta.append(["Exported by", user.email])
    meta.append(["Chain intact", audit.verify_chain(db).get("intact")])

    buffer = BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Audit_Trail.xlsx"'},
    )


AUDIT_COLUMNS = [
    ("ts_ist", "Timestamp (IST)"), ("entity_type", "Entity"),
    ("entity_id", "Entity ID"), ("action", "Action"), ("field", "Field"),
    ("old_value", "Old"), ("new_value", "New"), ("reason", "Reason"),
    ("user_email", "User"), ("ip", "IP"), ("hash", "Hash"), ("prev_hash", "Prev hash"),
]


def _serialise(row: AuditLog) -> dict:
    return {
        "id": row.id,
        "ts_utc": row.ts_utc,
        "ts_ist": row.ts_ist,
        "entity_type": row.entity_type,
        "entity_id": row.entity_id,
        "action": row.action,
        "field": row.field,
        "old_value": row.old_value,
        "new_value": row.new_value,
        "changes": row.changes,
        "reason": row.reason,
        "user_id": row.user_id,
        "user_email": row.user_email,
        "ip": row.ip,
        "user_agent": row.user_agent,
        "hash": row.hash,
        "prev_hash": row.prev_hash,
    }


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        import json

        return json.dumps(value, default=str)
    return str(value)
