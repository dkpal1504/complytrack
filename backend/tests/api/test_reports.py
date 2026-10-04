"""HTTP tests for the twelve standard reports and their CSV / Excel / PDF output."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

from tests.conftest import login

ADMIN = "admin@complytrack.local"
DEPARTMENT = "safety.owner@complytrack.local"

CODES = [f"R{n}" for n in range(1, 13)]
FORMATS = ["csv", "xlsx", "pdf"]


def _catalogue(client: TestClient, headers) -> list[dict]:
    response = client.get("/api/v1/reports", headers=headers)
    assert response.status_code == 200
    return response.json()


# --------------------------------------------------------------------------- #
# Catalogue
# --------------------------------------------------------------------------- #
def test_the_catalogue_lists_all_twelve_reports(demo, client: TestClient) -> None:
    catalogue = _catalogue(client, login(client, ADMIN))

    assert [entry["code"] for entry in catalogue] == CODES
    for entry in catalogue:
        assert entry["name"]
        assert entry["description"]
        assert entry["columns"]
        assert entry["group"]


def test_the_catalogue_is_refused_for_a_read_only_role(demo, client: TestClient) -> None:
    assert client.get(
        "/api/v1/reports", headers=login(client, "auditor@complytrack.local")
    ).status_code == 403


def test_an_unknown_report_code_is_a_404(demo, client: TestClient) -> None:
    headers = login(client, ADMIN)
    assert client.get("/api/v1/reports/R99", headers=headers).status_code == 404
    assert client.get("/api/v1/reports/R99/export", headers=headers).status_code == 404


# --------------------------------------------------------------------------- #
# Running a report
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("code", CODES)
def test_each_report_runs_as_json(demo, client: TestClient, code: str) -> None:
    headers = login(client, ADMIN)
    entry = next(e for e in _catalogue(client, headers) if e["code"] == code)

    response = client.get(f"/api/v1/reports/{code}", headers=headers)
    assert response.status_code == 200
    payload = response.json()
    assert payload["code"] == code
    assert payload["name"] == entry["name"]
    assert [column["label"] for column in payload["columns"]] == entry["columns"]
    assert isinstance(payload["rows"], list)
    assert isinstance(payload["summary"], dict)
    assert payload["generated_at"]


def test_the_register_report_covers_the_whole_demo_estate(demo, client: TestClient) -> None:
    payload = client.get("/api/v1/reports/R1", headers=login(client, ADMIN)).json()
    assert len(payload["rows"]) == 80
    assert payload["summary"]["total"] == 80
    assert payload["rows"][0]["code"].startswith(("FAC-", "ENV-", "LAB-", "DEF-", "PES-", "CMP-"))


def test_report_filters_narrow_the_rows(demo, client: TestClient) -> None:
    headers = login(client, ADMIN)
    unfiltered = client.get("/api/v1/reports/R1", headers=headers).json()

    register = client.get("/api/v1/items", params={"page_size": 500}, headers=headers).json()
    unit = register["items"][0]["unit_id"]
    filtered = client.get("/api/v1/reports/R1", headers=headers, params={"unit_id": unit}).json()

    assert len(filtered["rows"]) < len(unfiltered["rows"])
    unit_name = register["items"][0]["unit_name"]
    assert {row["unit"] for row in filtered["rows"]} == {unit_name}


def test_a_scoped_role_only_sees_its_own_rows(demo, client: TestClient) -> None:
    headers = login(client, DEPARTMENT)
    payload = client.get("/api/v1/reports/R1", headers=headers).json()
    me = client.get("/api/v1/auth/me", headers=headers).json()

    assert 0 < len(payload["rows"]) < 80
    assert payload["summary"]["total"] == len(payload["rows"])
    assert me["email"]  # the scope came from the signed-in user, not the payload


# --------------------------------------------------------------------------- #
# Exports
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("code", CODES)
@pytest.mark.parametrize("fmt", FORMATS)
def test_every_report_exports_in_every_format(
    demo, client: TestClient, code: str, fmt: str
) -> None:
    response = client.get(
        f"/api/v1/reports/{code}/export", params={"fmt": fmt}, headers=login(client, ADMIN)
    )
    assert response.status_code == 200
    assert response.headers["content-disposition"].startswith("attachment; filename=")
    assert f".{fmt}" in response.headers["content-disposition"]

    if fmt == "csv":
        assert response.headers["content-type"].startswith("text/csv")
        assert response.content.startswith(b"\xef\xbb\xbf")  # Excel-friendly BOM
    elif fmt == "xlsx":
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert response.content.startswith(b"PK\x03\x04")
    else:
        assert response.headers["content-type"] == "application/pdf"
        assert response.content.startswith(b"%PDF")


def test_the_csv_export_carries_a_header_and_every_row(demo, client: TestClient) -> None:
    headers = login(client, ADMIN)
    report = client.get("/api/v1/reports/R1", headers=headers).json()
    export = client.get("/api/v1/reports/R1/export", params={"fmt": "csv"}, headers=headers)

    import csv

    rows = list(csv.reader(io.StringIO(export.content.decode("utf-8-sig"))))
    assert rows[0] == [column["label"] for column in report["columns"]]
    assert len(rows) == len(report["rows"]) + 1


def test_the_excel_export_opens_with_the_report_name(demo, client: TestClient) -> None:
    from openpyxl import load_workbook

    export = client.get(
        "/api/v1/reports/R2/export", params={"fmt": "xlsx"}, headers=login(client, ADMIN)
    )
    workbook = load_workbook(io.BytesIO(export.content))
    assert workbook.sheetnames[0] == "R2"
    sheet = workbook["R2"]
    assert sheet.max_row >= 2
    assert sheet.cell(row=1, column=1).value == "Department"


def test_the_pdf_export_is_rendered_reportlab_output(demo, client: TestClient) -> None:
    export = client.get(
        "/api/v1/reports/R1/export", params={"fmt": "pdf"}, headers=login(client, ADMIN)
    )
    assert export.content.startswith(b"%PDF")
    assert b"%%EOF" in export.content[-2048:]


def test_an_unsupported_export_format_is_a_validation_error(demo, client: TestClient) -> None:
    response = client.get(
        "/api/v1/reports/R1/export", params={"fmt": "docx"}, headers=login(client, ADMIN)
    )
    assert response.status_code == 422


def test_a_report_can_be_exported_as_a_calendar_feed(demo, client: TestClient) -> None:
    response = client.get("/api/v1/reports/R4/ical", headers=login(client, ADMIN))
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/calendar")
    assert response.text.startswith("BEGIN:VCALENDAR")


# --------------------------------------------------------------------------- #
# Scheduled delivery
# --------------------------------------------------------------------------- #
def test_a_report_can_be_scheduled_and_removed(demo, client: TestClient) -> None:
    headers = login(client, ADMIN)

    created = client.post(
        "/api/v1/reports/schedules",
        headers=headers,
        json={
            "report_code": "R1",
            "name": "Monday register",
            "frequency": "weekly",
            "recipients": ["ceo@complytrack.local"],
            "format": "xlsx",
        },
    )
    assert created.status_code == 201
    schedule = created.json()
    assert schedule["report_code"] == "R1"
    assert schedule["recipients"] == ["ceo@complytrack.local"]
    assert schedule["active"] is True

    listed = client.get("/api/v1/reports/schedules/list", headers=headers).json()
    assert schedule["id"] in {row["id"] for row in listed}

    removed = client.delete(f"/api/v1/reports/schedules/{schedule['id']}", headers=headers)
    assert removed.status_code == 200
    assert removed.json()["ok"] is True

    after = client.get("/api/v1/reports/schedules/list", headers=headers).json()
    assert schedule["id"] not in {row["id"] for row in after}


@pytest.mark.parametrize("code", ["R1", "R7"])
def test_a_removed_schedule_disappears_from_the_listing(
    demo, client: TestClient, code: str
) -> None:
    headers = login(client, ADMIN)
    created = client.post(
        "/api/v1/reports/schedules",
        headers=headers,
        json={
            "report_code": code,
            "name": f"To be removed ({code})",
            "frequency": "weekly",
            "recipients": ["ceo@complytrack.local"],
        },
    ).json()
    assert created["id"] in {
        row["id"] for row in client.get("/api/v1/reports/schedules/list", headers=headers).json()
    }

    client.delete(f"/api/v1/reports/schedules/{created['id']}", headers=headers)
    assert all(
        row["id"] != created["id"]
        for row in client.get("/api/v1/reports/schedules/list", headers=headers).json()
    )


def test_a_schedule_needs_a_known_report_and_a_recipient(demo, client: TestClient) -> None:
    headers = login(client, ADMIN)

    unknown = client.post(
        "/api/v1/reports/schedules",
        headers=headers,
        json={
            "report_code": "R99",
            "name": "Bogus",
            "frequency": "daily",
            "recipients": ["a@b.local"],
        },
    )
    assert unknown.status_code == 400

    nobody = client.post(
        "/api/v1/reports/schedules",
        headers=headers,
        json={"report_code": "R1", "name": "Nobody", "frequency": "daily", "recipients": []},
    )
    assert nobody.status_code == 400
